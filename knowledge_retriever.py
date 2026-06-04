import os
import re
import numpy as np
from pathlib import Path
from typing import List, Dict, Optional, Tuple, Any
import logging

# 设置日志
logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

# 文本分割和向量化
try:
    from sentence_transformers import SentenceTransformer
except ImportError:
    logger.warning("sentence-transformers 未安装，文本嵌入功能将不可用")
    SentenceTransformer = None

# FAISS向量数据库
try:
    import faiss
except ImportError:
    logger.warning("faiss 未安装，将使用线性搜索替代")
    faiss = None

# PDF 处理
try:
    import pypdf
except ImportError:
    logger.warning("pypdf 未安装，PDF 处理功能将不可用")
    pypdf = None

# DOCX 处理
try:
    from docx import Document as DocxDocument
except ImportError:
    logger.warning("python-docx 未安装，DOCX 处理功能将不可用")
    DocxDocument = None


class DocumentKnowledgeRetriever:
    def __init__(self, knowledge_dir="knowledge_base", model_name='llm_models/all-MiniLM-L6-v2'):
        self.knowledge_dir = knowledge_dir
        self.model = None
        self.index = None
        self.documents: List[Dict] = []
        self.embeddings: Optional[np.ndarray] = None

        # 尝试加载模型
        self._load_model(model_name)

        # 自动加载知识库
        if os.path.exists(knowledge_dir):
            self._build_index()

    def _load_model(self, model_name):
        """加载嵌入模型"""
        if SentenceTransformer is None:
            logger.error("sentence-transformers 不可用，无法加载嵌入模型")
            return

        try:
            # 首先尝试本地路径
            local_path = f"models/{model_name}"
            if os.path.exists(local_path):
                logger.info(f"从本地加载嵌入模型: {local_path}")
                self.model = SentenceTransformer(local_path)
            else:
                # 尝试Hugging Face Hub
                try:
                    logger.info(f"从Hugging Face Hub加载嵌入模型: {model_name}")
                    self.model = SentenceTransformer(model_name)
                except Exception as e:
                    # 最后尝试sentence-transformers前缀
                    prefixed_name = f"sentence-transformers/{model_name}"
                    logger.info(f"尝试加载模型: {prefixed_name}")
                    self.model = SentenceTransformer(prefixed_name)
        except Exception as e:
            logger.error(f"加载嵌入模型失败: {str(e)}")
            self.model = None

    def _load_documents(self):
        """加载所有文档"""
        logger.info("加载知识库文档...")
        self.documents = []
        file_count = 0
        chunk_count = 0

        if not os.path.exists(self.knowledge_dir):
            logger.warning(f"知识库目录不存在: {self.knowledge_dir}")
            return

        for root, _, files in os.walk(self.knowledge_dir):
            for file in files:
                file_path = os.path.join(root, file)
                if file_path.endswith(('.pdf', '.docx', '.txt')):
                    file_count += 1
                    content = self._read_document(file_path)

                    if not content:
                        logger.warning(f"无法读取或空文件: {file_path}")
                        continue

                    # 分割文档
                    chunks = self._split_text(content)
                    chunk_count += len(chunks)

                    for i, chunk in enumerate(chunks):
                        self.documents.append({
                            "id": f"{Path(file).stem}-chunk-{i}",
                            "file": file,
                            "path": file_path,
                            "content": chunk,
                            "type": self._get_document_type(file)
                        })

        logger.info(f"已加载 {file_count} 个文件，共 {chunk_count} 个文档片段")

    def _get_document_type(self, filename):
        """根据文件名判断文档类型"""
        filename_lower = filename.lower()
        if 'attack' in filename_lower or 'mitre' in filename_lower:
            return "attack_framework"
        elif 'cve' in filename_lower or 'vulnerability' in filename_lower:
            return "threat_intel"
        elif 'case' in filename_lower or 'incident' in filename_lower:
            return "threat_case"
        elif 'response' in filename_lower or 'playbook' in filename_lower:
            return "response_playbook"
        return "other"

    def _read_document(self, file_path):
        """读取不同格式的文档内容"""
        try:
            if file_path.endswith('.pdf'):
                return self._read_pdf(file_path)
            elif file_path.endswith('.docx'):
                return self._read_docx(file_path)
            elif file_path.endswith('.txt'):
                return self._read_text(file_path)
            else:
                logger.warning(f"不支持的文档格式: {file_path}")
                return ""
        except Exception as e:
            logger.error(f"读取文档失败: {file_path} - {str(e)}")
            return ""

    def _read_pdf(self, file_path):
        """读取PDF文件 - 使用pypdf库"""
        if pypdf is None:
            logger.warning("pypdf未安装，无法读取PDF")
            return ""

        content = []
        try:
            with open(file_path, 'rb') as f:
                reader = pypdf.PdfReader(f)
                for page in reader.pages:
                    text = page.extract_text()
                    if text:
                        content.append(text)
            return "\n".join(content)
        except Exception as e:
            logger.error(f"读取PDF失败: {file_path} - {str(e)}")
            return ""

    def _read_docx(self, file_path):
        """读取DOCX文件"""
        if DocxDocument is None:
            logger.warning("python-docx未安装，无法读取DOCX")
            return ""

        try:
            doc = DocxDocument(file_path)
            full_text = []

            # 处理段落
            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text)

            # 处理表格
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for para in cell.paragraphs:
                            if para.text.strip():
                                full_text.append(para.text)

            return "\n".join(full_text)
        except Exception as e:
            logger.error(f"读取DOCX失败: {file_path} - {str(e)}")
            return ""

    def _read_text(self, file_path):
        """读取文本文件"""
        try:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read()
        except Exception as e:
            logger.error(f"读取文本文件失败: {file_path} - {str(e)}")
            return ""

    def _split_text(self, text, chunk_size=500, chunk_overlap=50):
        """自定义文本分割函数"""
        if not text:
            return []

        chunks = []
        start = 0
        text_length = len(text)

        while start < text_length:
            end = min(start + chunk_size, text_length)

            # 尝试在句子边界处分割
            if end < text_length:
                # 查找最近的句子结束点
                for boundary in ['.', '!', '?', '\n\n', '\n']:
                    pos = text.rfind(boundary, start, end)
                    if pos != -1 and pos > start + chunk_size // 2:
                        end = pos + len(boundary)
                        break

            chunks.append(text[start:end].strip())

            # 计算下一个开始位置（考虑重叠）
            start = max(start + 1, end - chunk_overlap)

        return chunks

    def _build_index(self):
        """构建向量索引"""
        self._load_documents()
        if not self.documents:
            logger.warning("知识库为空")
            return

        if self.model is None:
            logger.error("嵌入模型未加载，无法构建索引")
            return

        logger.info("构建向量索引...")
        # 生成嵌入向量
        contents = [doc["content"] for doc in self.documents]
        self.embeddings = self.model.encode(
            contents,
            show_progress_bar=True,
            batch_size=32,
            convert_to_numpy=True
        )

        # 创建FAISS索引（如果可用）
        if faiss is not None:
            try:
                dimension = self.embeddings.shape[1]
                self.index = faiss.IndexFlatL2(dimension)
                self.index.add(self.embeddings.astype(np.float32))
                logger.info("FAISS向量索引构建完成")
            except Exception as e:
                logger.error(f"构建FAISS索引失败: {str(e)}")
                self.index = None
        else:
            logger.info("FAISS不可用，将使用线性搜索")
            self.index = None

    def _linear_search(self, query_embed, top_k) -> Tuple[np.ndarray, np.ndarray]:
        """当FAISS不可用时使用线性搜索"""
        if self.embeddings is None or len(self.documents) == 0:
            return np.array([]), np.array([])

        # 计算所有文档与查询的余弦相似度
        query_norm = np.linalg.norm(query_embed)
        doc_norms = np.linalg.norm(self.embeddings, axis=1)
        similarities = np.dot(self.embeddings, query_embed.T).flatten() / (doc_norms * query_norm + 1e-9)

        # 获取top_k结果
        indices = np.argsort(similarities)[::-1][:top_k]
        distances = 1 - similarities[indices]  # 转换为距离（越小越好）

        return distances.reshape(1, -1), indices.reshape(1, -1)

    def retrieve(self, query, top_k=5, doc_type=None) -> List[Dict[str, Any]]:
        """检索知识库"""
        if not self.documents:
            logger.warning("知识库为空")
            return []

        if self.model is None:
            logger.error("嵌入模型未加载，无法检索")
            return []

        # 生成查询向量
        query_embed = self.model.encode([query])[0].reshape(1, -1)

        # 执行搜索
        if self.index is not None:
            try:
                distances, indices = self.index.search(
                    query_embed.astype(np.float32),
                    top_k * 3  # 初始获取更多结果以便过滤
                )
            except Exception as e:
                logger.error(f"FAISS搜索失败: {str(e)}，回退到线性搜索")
                distances, indices = self._linear_search(query_embed, top_k * 3)
        else:
            distances, indices = self._linear_search(query_embed, top_k * 3)

        # 过滤和排序结果
        results = []
        seen_docs = set()

        for i, idx in enumerate(indices[0]):
            if idx < 0 or idx >= len(self.documents):
                continue

            doc = self.documents[idx]

            # 类型过滤
            if doc_type and doc["type"] != doc_type:
                continue

            # 去重
            if doc["id"] in seen_docs:
                continue
            seen_docs.add(doc["id"])

            # 添加结果
            results.append({
                "id": doc["id"],
                "file": doc["file"],
                "type": doc["type"],
                "content": doc["content"],
                "score": float(distances[0][i]) if distances.size > 0 else 0.0
            })

            # 达到所需数量
            if len(results) >= top_k:
                break

        # 按相关度排序（分数越低越好）
        results.sort(key=lambda x: x["score"])
        return results

    def get_attack_chain(self, technique_id):
        """从文档中提取攻击链信息"""
        # 检索相关文档
        results = self.retrieve(technique_id, top_k=3, doc_type="attack_framework")
        if not results:
            return None

        # 从文档内容提取攻击链
        attack_chain = []

        for result in results:
            content = result["content"]

            # 使用正则表达式提取攻击链信息
            patterns = [
                r"攻击链[:：]\s*(.*)",
                r"kill chain[:：]\s*(.*)",
                r"攻击阶段[:：]\s*([^\n]+)",
                r"阶段\s*\d+[:：]\s*([^\n]+)"
            ]

            for pattern in patterns:
                matches = re.findall(pattern, content)
                for match in matches:
                    if match.strip():
                        attack_chain.append(match.strip())

        return attack_chain if attack_chain else ["未找到攻击链信息"]

    def get_incident_response(self, threat_type):
        """获取事件响应建议"""
        results = self.retrieve(threat_type, top_k=3, doc_type="response_playbook")
        if not results:
            return []

        response_steps = []

        for result in results:
            content = result["content"]

            # 提取响应建议
            patterns = [
                r"响应步骤[:：]\s*([^\n]+)",
                r"建议\s*\d+[:：]\s*([^\n]+)",
                r"立即行动[:：]\s*([^\n]+)",
                r"long-term[:：]\s*([^\n]+)",
                r"步骤\s*\d+[:：]\s*([^\n]+)"
            ]

            for pattern in patterns:
                matches = re.findall(pattern, content)
                for match in matches:
                    if match.strip():
                        response_steps.append(match.strip())

        return response_steps if response_steps else ["未找到响应建议"]

    def add_document(self, file_path):
        """添加单个文档到知识库"""
        if not os.path.exists(file_path):
            logger.error(f"文件不存在: {file_path}")
            return False

        # 确保知识目录存在
        os.makedirs(self.knowledge_dir, exist_ok=True)

        # 复制文件到知识库
        import shutil
        try:
            dest_path = os.path.join(self.knowledge_dir, os.path.basename(file_path))
            shutil.copy2(file_path, dest_path)
            logger.info(f"已添加文档: {dest_path}")

            # 重新构建索引
            self._build_index()
            return True
        except Exception as e:
            logger.error(f"添加文档失败: {str(e)}")
            return False

    def search(self, query, top_k=5):
        """通用搜索接口"""
        return self.retrieve(query, top_k)


if __name__ == "__main__":
    # 创建知识库目录
    os.makedirs("knowledge_base", exist_ok=True)

    # 生成测试文件
    test_files = {
        "MITRE_ATT&CK_Enterprise.txt": """MITRE ATT&CK框架企业版技术列表\n\nT1059 - 命令行界面\n攻击链:\n阶段1: 初始访问...""",
        "CVE-2023-1234.txt": """CVE-2023-1234漏洞通告\n\n漏洞名称: Apache Log4j远程代码执行...""",
        "Incident_Response_Playbook.txt": """网络安全事件响应手册\n\n勒索软件响应流程:\n[立即行动]\n1. 隔离感染主机..."""
    }

    for filename, content in test_files.items():
        with open(f"knowledge_base/{filename}", "w", encoding='utf-8') as f:
            f.write(content)

    # 初始化知识检索器
    retriever = DocumentKnowledgeRetriever(knowledge_dir="knowledge_base")

    # 测试功能
    print("攻击链信息:", retriever.get_attack_chain("T1059"))
    print("响应建议:", retriever.get_incident_response("勒索软件"))
    print("通用搜索:", retriever.search("Log4j漏洞", top_k=2))